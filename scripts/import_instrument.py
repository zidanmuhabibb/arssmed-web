"""
Impor instrumen tes diagnostik four-tier dari berkas .docx peneliti.

    pip install python-docx
    python3 scripts/import_instrument.py "Instrumen_Tes_Diagnostik_Four-Tier_Tata_Surya.docx"

Menulis (menimpa):
  data/items.json                    20 butir + kunci + metadata kisi-kisi
  data/rule-sets/pedoman-v1.json     Tabel D.2 (16 baris) → aturan per_tier
  tests/fixtures/pedoman-examples.json  Contoh D.3 (dipakai unit test)

Skrip berhenti dengan galat bila struktur dokumen tidak sesuai harapan
(jumlah butir ≠ 20, opsi ≠ A–D, kunci rinci ≠ kunci ringkasan, dll.),
sehingga perubahan dokumen tidak diam-diam menghasilkan data salah.
"""
import json
import re
import sys
from pathlib import Path

import docx

ROOT = Path(__file__).resolve().parent.parent
LETTERS = ["A", "B", "C", "D"]

DOMAINS = [
    # (id, label di dokumen, nomor butir) — kisi-kisi bagian A dan rekap D.5
    ("benda_langit", "Karakteristik benda langit", range(1, 6)),
    ("planet", "Susunan dan karakteristik planet", range(6, 11)),
    ("rotasi_revolusi", "Rotasi dan revolusi Bumi", range(11, 16)),
    ("gerhana", "Fenomena astronomi (gerhana)", range(16, 21)),
]
DOMAIN_OF = {n: d for d, _, nums in DOMAINS for n in nums}

CATEGORY_CODE = {"SC", "LC", "M", "LK", "E"}


def fail(msg):
    sys.exit(f"GAGAL: {msg}")


def norm(text):
    text = text.replace(" ", " ").replace("\t", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def parse_items(doc):
    paras = [norm(p.text) for p in doc.paragraphs]
    items = {}
    i = 0
    while i < len(paras):
        m = re.match(r"^(\d{1,2})\. (.+)$", paras[i])
        if m and i + 1 < len(paras) and paras[i + 1].startswith("Tier 1"):
            no = int(m.group(1))
            stem = m.group(2)
            j = i + 2
            t1, t3 = [], []
            while not paras[j].startswith("Tier 2"):
                t1.append(paras[j]); j += 1
            j += 1
            if not paras[j].startswith("Tier 3"):
                fail(f"butir {no}: 'Tier 3' tidak ditemukan setelah keyakinan tier 2")
            j += 1
            while not paras[j].startswith("Tier 4"):
                t3.append(paras[j]); j += 1
            items[no] = {"stem": stem, "t1": options(no, "tier 1", t1), "t3": options(no, "tier 3", t3),
                         "conf2": paras[i + 2 + len(t1)], "conf4": paras[j]}
            i = j + 1
        else:
            i += 1
    if sorted(items) != list(range(1, 21)):
        fail(f"butir yang ditemukan: {sorted(items)} (harus 1–20)")
    return items


def options(no, tier, lines):
    out = []
    for line in lines:
        m = re.match(r"^([A-D])\.\s+(.+)$", line)
        if not m:
            fail(f"butir {no} {tier}: baris opsi tidak dikenali: {line!r}")
        out.append({"key": m.group(1), "text": m.group(2)})
    if [o["key"] for o in out] != LETTERS:
        fail(f"butir {no} {tier}: opsi harus A–D, ditemukan {[o['key'] for o in out]}")
    return out


def table_rows(table):
    return [[norm(c.text) for c in r.cells] for r in table.rows]


def find_table(doc, header_start):
    for t in doc.tables:
        if norm(t.rows[0].cells[0].text).startswith(header_start):
            return table_rows(t)
    fail(f"tabel dengan kolom pertama '{header_start}' tidak ditemukan")


def main(path):
    doc = docx.Document(path)
    items = parse_items(doc)

    # Keyakinan: harus "Yakin" / "Tidak yakin" (dua level)
    for no, it in items.items():
        for k in ("conf2", "conf4"):
            if "Yakin" not in it[k] or "Tidak yakin" not in it[k]:
                fail(f"butir {no}: skala keyakinan tak terduga: {it[k]!r}")

    # Kisi-kisi (A)
    kisi = {}
    for row in find_table(doc, "Domain konsep")[1:]:
        no = int(row[1]); kisi[no] = {"indicator": row[2], "alternative_conceptions": row[3]}

    # Kunci rinci (C)
    key = {}
    detail = None
    for t in doc.tables:
        rows = table_rows(t)
        if rows[0][:3] == ["No", "Tier 1", "Tier 3"]:
            detail = rows
    if detail is None:
        fail("tabel kunci jawaban rinci tidak ditemukan")
    for row in detail[1:]:
        key[int(row[0])] = {"t1": row[1], "t3": row[2], "scientific_concept": row[3], "distractors": row[4]}

    # Kunci ringkasan — harus sama dengan kunci rinci
    summary = None
    for t in doc.tables:
        rows = table_rows(t)
        if rows[0][0] == "No" and len(rows[0]) == 11:
            summary = rows
    if summary is None:
        fail("tabel ringkasan kunci tidak ditemukan")
    sum_key = {}
    for b in (0, 3):
        nums = [int(x) for x in summary[b][1:]]
        for n, a, r in zip(nums, summary[b + 1][1:], summary[b + 2][1:]):
            sum_key[n] = (a, r)
    for n in range(1, 21):
        if sum_key[n] != (key[n]["t1"], key[n]["t3"]):
            fail(f"butir {n}: kunci rinci {key[n]['t1']}/{key[n]['t3']} ≠ ringkasan {sum_key[n]}")

    # Tabel keputusan D.2
    d2 = None
    for t in doc.tables:
        rows = table_rows(t)
        if rows[0][0] == "No" and rows[0][1].startswith("Tier 1 (jawaban)"):
            d2 = rows
    if d2 is None or len(d2) != 17:
        fail("tabel keputusan D.2 (16 baris) tidak ditemukan")
    rules = []
    for row in d2[1:]:
        cat = row[5].split("–")[0].strip()
        if cat not in CATEGORY_CODE:
            fail(f"kategori tak dikenal di D.2: {row[5]!r}")
        rules.append({"no": int(row[0]), "A": row[1] == "B", "CA": row[2] == "Y",
                      "R": row[3] == "B", "CR": row[4] == "Y", "category": cat})

    # Contoh D.3 (butir 7)
    d3 = None
    for t in doc.tables:
        rows = table_rows(t)
        if rows[0][0].startswith("Respons (T1, T2, T3, T4)"):
            d3 = rows
    if d3 is None:
        fail("tabel contoh D.3 tidak ditemukan")
    examples = []
    for row in d3[1:]:
        t1, t2, t3, t4 = [x.strip() for x in row[0].split(",")]
        examples.append({"item_order": 7, "tier1": t1, "confidenceA": t2, "reason": t3,
                         "confidenceR": t4, "interpretation": row[1], "category": row[2]})

    source = {
        "document": Path(path).name,
        "author": "Zidan Muhabib",
        "institution": "Program Studi Magister Pendidikan Dasar, Universitas Muhammadiyah Purwokerto",
        "year": 2026,
        "status": "draft_needs_validation",
        "status_note": ("Dokumen menyatakan butir, kunci, dan aturan pengkategorian adalah hasil penyusunan baru, "
                        "bukan salinan instrumen asli artikel. CVI 0,87 dan KR-20 0,79 berlaku untuk instrumen asli; "
                        "validasi isi oleh ahli dan uji reliabilitas perlu diulang sebelum dipakai."),
    }

    out_items = []
    for n in range(1, 21):
        it = items[n]
        domain = DOMAIN_OF[n]
        out_items.append({
            "item_order": n,
            "item_code": f"B{n:02d}",
            "content": {
                "stem": it["stem"],
                "stem_image": None,
                "format": "four_tier_standard",
                "tier1": {"options": it["t1"], "correct": key[n]["t1"], "fixed_order": False},
                "reason": {"options": [{**o, "maps_to_misconception": None} for o in it["t3"]],
                           "correct": key[n]["t3"]},
                "confidence": {"levels": ["Yakin", "Tidak yakin"], "threshold_index": 0},
                "concept_domain": domain,
                "report_domain": domain,
                "misconception_target": f"MK-B{n:02d}",
            },
            "meta": {
                "indicator": kisi[n]["indicator"],
                "alternative_conceptions": kisi[n]["alternative_conceptions"],
                "scientific_concept": key[n]["scientific_concept"],
                "distractor_notes": key[n]["distractors"],
            },
        })

    items_file = {
        "test_name": "Tes diagnostik four-tier tata surya",
        "test_version": 1,
        "source": source,
        "default_rule_set_id": "pedoman-v1",
        "domains": [{"id": d, "label": label, "item_orders": list(nums)} for d, label, nums in DOMAINS],
        "items": out_items,
    }
    rule_set = {
        "rule_set_id": "pedoman-v1",
        "kind": "per_tier",
        "source": f"{source['document']}, bagian D.2 Tabel Keputusan Kategori. Tier kosong → E (catatan di bawah tabel D.2).",
        "incomplete_category": "E",
        "rules": rules,
    }

    def write(rel, obj):
        p = ROOT / rel
        p.write_text(json.dumps(obj, ensure_ascii=False, indent=2) + "\n", encoding="utf8")
        print(f"ditulis: {rel}")

    write("data/items.json", items_file)
    write("data/rule-sets/pedoman-v1.json", rule_set)
    write("tests/fixtures/pedoman-examples.json", {"source": source["document"] + " bagian D.3", "examples": examples})


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
